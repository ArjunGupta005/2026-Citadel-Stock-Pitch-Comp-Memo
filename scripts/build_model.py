"""Build model/TeamName.xlsx: live-formula CPRT short model (KPI -> EPS -> justified multiple -> distribution).

Mirrors scripts/scenario_engine.cprt_quarters and scripts/cprt_framework_b exactly (verified by verify_model.py).
Blinded: no team, school or names.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import cprt_framework_b as fb
import scenario_engine as se
from common import ROOT

NAVY = "1F3864"
HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor=NAVY)
INPUT = Font(color="0000FF")            # blue = hard-coded input
CALC = Font(color="000000")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
PCT, NUM2, NUM1, USD = "0.0%", "0.00", "0.0", '"$"#,##0.00'
wb = Workbook()


def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row, col + i, lab)
        c.font, c.fill = HDR, FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")


def put(ws, ref, v, fmt=None, font=None):
    ws[ref] = v
    if fmt:
        ws[ref].number_format = fmt
    if font:
        ws[ref].font = font


# ---------------------------------------------------------------- Inputs
I = wb.active
I.title = "Inputs"
I["A1"] = "Inputs (blue = hard-coded; sources on 'Sources' tab)"
I["A1"].font = Font(bold=True, size=13)
rows = [
    ("Share price (2026-10-01 close)", 27.29, USD), ("Diluted shares (m)", 940.0, NUM1), ("Tax rate", 0.195, PCT),
    ("Cash + securities pre-ACV ($bn)", 4.7, NUM2), ("Cash at FY27 year-end post-ACV ($bn)", 3.5, NUM2),
    ("10y UST (risk-free)", 0.0524, "0.00%"), ("Equity beta", 0.90, NUM2), ("Equity risk premium", 0.05, PCT),
    ("Terminal growth", 0.03, PCT), ("FY26 interest income ($m)", 181.9, NUM1), ("US insurance share of US units", 0.76, PCT),
    ("Fixed share of US facility cost", 0.55, PCT), ("US non-insurance unit growth", 0.03, PCT),
    ("US vehicle GP / G&A growth", 0.05, PCT), ("International op. income growth", 0.10, PCT),
    ("Interest per quarter pre-ACV close ($m)", 38.0, NUM1),
]
for i, (lab, v, f) in enumerate(rows, start=3):
    I.cell(i, 1, lab)
    put(I, f"B{i}", v, f, INPUT)
NAMES = {lab: f"Inputs!$B${i}" for i, (lab, _, _) in enumerate(rows, start=3)}
PRICE, SH, TAX = NAMES["Share price (2026-10-01 close)"], NAMES["Diluted shares (m)"], NAMES["Tax rate"]
CASH0, CASH1 = NAMES["Cash + securities pre-ACV ($bn)"], NAMES["Cash at FY27 year-end post-ACV ($bn)"]
RF, BETA, ERP, GT = NAMES["10y UST (risk-free)"], NAMES["Equity beta"], NAMES["Equity risk premium"], NAMES["Terminal growth"]
INT26, SINS, FIX = NAMES["FY26 interest income ($m)"], NAMES["US insurance share of US units"], NAMES["Fixed share of US facility cost"]
NONINS, GOTH, GINT, INTQ1 = (NAMES["US non-insurance unit growth"], NAMES["US vehicle GP / G&A growth"],
                             NAMES["International op. income growth"], NAMES["Interest per quarter pre-ACV close ($m)"])
r0 = 3 + len(rows) + 1
I.cell(r0, 1, "Cost of equity (= rf + beta × ERP)").font = BOLD
I[f"B{r0}"] = f"={RF}+{BETA}*{ERP}"
I[f"B{r0}"].number_format = "0.00%"
COE = f"Inputs!$B${r0}"
# FY26 quarterly segment actuals
r1 = r0 + 2
I.cell(r1, 1, "FY26 quarterly segment actuals ($m) — 10-Qs / 10-K FY26 (Q4 = FY − 9M)").font = BOLD
header(I, r1 + 1, ["Line", "Q1 (Aug-Oct)", "Q2 (Nov-Jan)", "Q3 (Feb-Apr)", "Q4 (May-Jul)"])
seg = [("US service revenue", "us_service"), ("US vehicle sales", "us_vehicle"), ("US facility operations", "us_facility"),
       ("US cost of vehicle sales", "us_cogs_vehicle"), ("US G&A", "us_ga"), ("International op. income", "intl_oi"),
       ("Interest income", "interest")]
SEG = {}
for j, (lab, key) in enumerate(seg):
    rr = r1 + 2 + j
    I.cell(rr, 1, lab)
    for q in range(4):
        put(I, f"{get_column_letter(2 + q)}{rr}", se.CPRT_FY26[key][q], NUM1, INPUT)
    SEG[key] = rr
rr = r1 + 2 + len(seg)
I.cell(rr, 1, "Share of Progressive run-rate loss already in FY26 quarter")
for q in range(4):
    put(I, f"{get_column_letter(2 + q)}{rr}", se.CPRT_LOSS_IN_BASE[q], PCT, INPUT)
SEG["loss_in_base"] = rr
I.column_dimensions["A"].width = 58
for c in "BCDE":
    I.column_dimensions[c].width = 14

# ---------------------------------------------------------------- Scenarios
S = wb.create_sheet("Scenarios")
S["A1"] = "Scenario worlds (economic worlds: what the market thinks → KPIs → numbers → multiple)"
S["A1"].font = Font(bold=True, size=13)
cols = ["Parameter", "W1 Share loss deepens", "W2 FY27 trough → FY28 normalises", "W3 Cycle turns + ACV synergies",
        "W4 Street right (share recapture)"]
header(S, 3, cols)
worlds = [se.CPRT_WORLDS["bear"], se.CPRT_WORLDS["base"], se.CPRT_WORLDS["bull"], fb.W4]
extra = [(0.20, 0.04, 0.03, 0.15), (0.45, 0.12, 0.06, 0.20), (0.20, 0.18, 0.09, 0.25), (0.15, 0.15, 0.10, 0.25)]
params = [("Underlying US insurance unit growth (ex lost customer)", "u", PCT),
          ("Progressive run-rate loss (pts of US insurance units)", "loss", PCT), ("Revenue per unit growth", "rpu", PCT),
          ("Fixed facility cost growth (investment)", "fixed", PCT), ("ACV amortization per quarter ($m)", "acv_amort", NUM1),
          ("ACV operating income per quarter ($m)", "acv_op", NUM1), ("Interest income per quarter post-close ($m)", "int_q", NUM1)]
P = {}
for i, (lab, key, f) in enumerate(params, start=4):
    S.cell(i, 1, lab)
    for j, w in enumerate(worlds):
        put(S, f"{get_column_letter(2 + j)}{i}", w[key], f, INPUT)
    P[key] = i
for k, (lab, f) in enumerate([("Probability", PCT), ("FY28 EPS growth", PCT), ("5-yr EPS growth after FY28", PCT),
                              ("Incremental ROIC", PCT)]):
    i = 4 + len(params) + k
    S.cell(i, 1, lab)
    for j in range(4):
        put(S, f"{get_column_letter(2 + j)}{i}", extra[j][k], f, INPUT)
    P[["prob", "g28", "g5", "roic"][k]] = i
i = P["roic"] + 1
S.cell(i, 1, "Probability check (must = 100%)").font = BOLD
S[f"B{i}"] = f"=SUM(B{P['prob']}:E{P['prob']})"
S[f"B{i}"].number_format = PCT
nar = P["roic"] + 3
S.cell(nar, 1, "Narrative / evidence for each world").font = BOLD
texts = ["Share loss spreads beyond Progressive; Manheim negative cuts ASP; OpEx/car keeps rising; ACV dilutes ROIC.",
         "Progressive loss runs at full rate Q1-Q3 (RBA: all 50 states, 90-day transition); laps in Q4; cost actions; ACV breakeven → accretive.",
         "Record total-loss frequency; ex-customer US insurance +3%; RPU +5%; $209m ACV synergies (HSBC).",
         "JPM proprietary data: share gains with a carrier offset most of the loss; Street FY27 roughly delivered."]
for j, t in enumerate(texts):
    c = S.cell(nar + 1, 2 + j, t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
S.row_dimensions[nar + 1].height = 95
S.column_dimensions["A"].width = 52
for c in "BCDE":
    S.column_dimensions[c].width = 26

# ---------------------------------------------------------------- EPS build
E = wb.create_sheet("EPS_Build")
E["A1"] = "FY27 (Aug-26 → Jul-27) quarterly EPS build by world — GAAP diluted; ACV consolidated from Q2 (close ~mid-Oct-26)"
E["A1"].font = Font(bold=True, size=13)
lines = ["Incremental Progressive loss vs FY26 quarter", "US insurance units y/y", "US units y/y", "US service revenue ($m)",
         "US vehicle gross profit ($m)", "US facility operations ($m)", "US G&A ($m)", "US operating income ($m)",
         "International operating income ($m)", "ACV pre-tax (op. income − amortization) ($m)", "Interest income ($m)",
         "Pre-tax income ($m)", "Diluted EPS", "EPS ex-ACV (no ACV P&L, cash not spent)"]
EPS_CELLS, EPSX_CELLS = {}, {}
for j in range(4):
    top = 3 + j * 18
    wc = get_column_letter(2 + j)  # world column on Scenarios
    E.cell(top, 1, cols[1 + j]).font = BOLD
    header(E, top + 1, ["Line", "Q1", "Q2", "Q3", "Q4", "FY27"])
    for k, lab in enumerate(lines):
        E.cell(top + 2 + k, 1, lab)
    for q in range(4):
        c = get_column_letter(2 + q)
        R = {lab: f"{c}{top + 2 + k}" for k, lab in enumerate(lines)}
        S_ = lambda key: f"Scenarios!${wc}${P[key]}"
        segc = lambda key: f"Inputs!{c}${SEG[key]}"
        f = {
            lines[0]: f"={S_('loss')}*(1-{segc('loss_in_base')})",
            lines[1]: f"={S_('u')}+{R[lines[0]]}",
            lines[2]: f"={SINS}*{R[lines[1]]}+(1-{SINS})*{NONINS}",
            lines[3]: f"={segc('us_service')}*(1+{R[lines[2]]})*(1+{S_('rpu')})",
            lines[4]: f"=({segc('us_vehicle')}-{segc('us_cogs_vehicle')})*(1+{GOTH})",
            lines[5]: f"={segc('us_facility')}*{FIX}*(1+{S_('fixed')})+{segc('us_facility')}*(1-{FIX})*(1+{R[lines[2]]})",
            lines[6]: f"={segc('us_ga')}*(1+{GOTH})",
            lines[7]: f"={R[lines[3]]}+{R[lines[4]]}-{R[lines[5]]}-{R[lines[6]]}",
            lines[8]: f"={segc('intl_oi')}*(1+{GINT})",
            lines[9]: "=0" if q == 0 else f"={S_('acv_op')}-{S_('acv_amort')}",
            lines[10]: f"={INTQ1}" if q == 0 else f"={S_('int_q')}",
            lines[11]: f"={R[lines[7]]}+{R[lines[8]]}+{R[lines[9]]}+{R[lines[10]]}",
            lines[12]: f"={R[lines[11]]}*(1-{TAX})/{SH}",
            lines[13]: f"=({R[lines[7]]}+{R[lines[8]]}+{INTQ1})*(1-{TAX})/{SH}",
        }
        for lab, formula in f.items():
            E[R[lab]] = formula
            E[R[lab]].number_format = PCT if lab in lines[:3] else (NUM2 if "EPS" in lab else NUM1)
    for k in (3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13):
        rr = top + 2 + k
        E[f"F{rr}"] = f"=SUM(B{rr}:E{rr})"
        E[f"F{rr}"].number_format = NUM2 if "EPS" in lines[k] else NUM1
        E[f"F{rr}"].font = BOLD
    EPS_CELLS[j] = f"EPS_Build!$F${top + 2 + 12}"
    EPSX_CELLS[j] = f"EPS_Build!$F${top + 2 + 13}"
E.column_dimensions["A"].width = 48
for c in "BCDEF":
    E.column_dimensions[c].width = 13

# ---------------------------------------------------------------- Valuation
V = wb.create_sheet("Valuation")
V["A1"] = "Valuation: justified ex-cash P/E from growth, ROIC and cost of equity (no free multiples) → probability-weighted value"
V["A1"].font = Font(bold=True, size=13)
header(V, 3, ["Line"] + cols[1:])
vl = ["Probability", "FY27 EPS (GAAP, incl. ACV)", "FY27 EPS ex-ACV", "FY28 EPS", "After-tax interest per share (FY28)",
      "FY28 ex-interest EPS", "5-yr growth g", "Incremental ROIC", "Cost of equity", "Terminal growth",
      "PV of 5-yr distributions per $1 NTM EPS", "PV of terminal per $1 NTM EPS", "Justified ex-cash P/E",
      "Cash per share at horizon", "Value per share (Sep-27)", "Return vs price", "Implied total P/E on FY28"]
VR = {lab: 4 + k for k, lab in enumerate(vl)}
for lab, rr in VR.items():
    V.cell(rr, 1, lab)
for j in range(4):
    c = get_column_letter(2 + j)
    wc = get_column_letter(2 + j)
    ref = lambda lab: f"{c}{VR[lab]}"
    g, ro, r = ref("5-yr growth g"), ref("Incremental ROIC"), ref("Cost of equity")
    pv5 = "+".join(f"(1+{g})^{t - 1}*(1-{g}/{ro})/(1+{r})^{t}" for t in range(1, 6))
    fm = {
        "Probability": f"=Scenarios!{wc}{P['prob']}",
        "FY27 EPS (GAAP, incl. ACV)": f"={EPS_CELLS[j]}",
        "FY27 EPS ex-ACV": f"={EPSX_CELLS[j]}",
        "FY28 EPS": f"={ref('FY27 EPS (GAAP, incl. ACV)')}*(1+Scenarios!{wc}{P['g28']})",
        "After-tax interest per share (FY28)": f"=Scenarios!{wc}{P['int_q']}*4*(1-{TAX})/{SH}",
        "FY28 ex-interest EPS": f"={ref('FY28 EPS')}-{ref('After-tax interest per share (FY28)')}",
        "5-yr growth g": f"=Scenarios!{wc}{P['g5']}",
        "Incremental ROIC": f"=Scenarios!{wc}{P['roic']}",
        "Cost of equity": f"={COE}",
        "Terminal growth": f"={GT}",
        "PV of 5-yr distributions per $1 NTM EPS": f"={pv5}",
        "PV of terminal per $1 NTM EPS": f"=(1+{g})^5*(1-{ref('Terminal growth')}/{ro})/({r}-{ref('Terminal growth')})/(1+{r})^5",
        "Justified ex-cash P/E": f"={ref('PV of 5-yr distributions per $1 NTM EPS')}+{ref('PV of terminal per $1 NTM EPS')}",
        "Cash per share at horizon": f"={CASH1}*1000/{SH}",
        "Value per share (Sep-27)": f"={ref('Justified ex-cash P/E')}*{ref('FY28 ex-interest EPS')}+{ref('Cash per share at horizon')}",
        "Return vs price": f"={ref('Value per share (Sep-27)')}/{PRICE}-1",
        "Implied total P/E on FY28": f"={ref('Value per share (Sep-27)')}/{ref('FY28 EPS')}",
    }
    for lab, formula in fm.items():
        V[ref(lab)] = formula
        V[ref(lab)].number_format = (PCT if lab in ("Probability", "5-yr growth g", "Incremental ROIC", "Cost of equity",
                                                    "Terminal growth", "Return vs price")
                                     else USD if lab in ("Value per share (Sep-27)", "Cash per share at horizon")
                                     else "0.0x" if "P/E" in lab else NUM2)
s0 = 4 + len(vl) + 1
summ = [("Probability-weighted value per share", f"=SUMPRODUCT(B{VR['Probability']}:E{VR['Probability']},B{VR['Value per share (Sep-27)']}:E{VR['Value per share (Sep-27)']})", USD),
        ("Probability-weighted return (stock)", f"=B{s0}/{PRICE}-1", PCT),
        ("Expected return for a SHORT", f"=-B{s0 + 1}", PCT),
        ("Base-case (W2) price target", f"=C{VR['Value per share (Sep-27)']}", USD),
        ("Maximum plausible loss for a short (worst world)", f"=MAX(B{VR['Return vs price']}:E{VR['Return vs price']})", PCT),
        ("Market ex-cash, ex-interest NTM P/E on S&P FY27 $1.62", f"=({PRICE}-{CASH0}*1000/{SH})/(Consensus!C5-{INT26}*0.6*(1-{TAX})/{SH})", "0.0x")]
for k, (lab, formula, f) in enumerate(summ):
    V.cell(s0 + k, 1, lab).font = BOLD
    V[f"B{s0 + k}"] = formula
    V[f"B{s0 + k}"].number_format = f
    V[f"B{s0 + k}"].font = BOLD
VSUM = {lab: f"Valuation!$B${s0 + k}" for k, (lab, _, _) in enumerate(summ)}
e0 = s0 + len(summ) + 1
V.cell(e0, 1, "EPS-only check: today's market multiple held, only EPS changes (does the short need multiple compression?)").font = BOLD
for j in range(4):
    c = get_column_letter(2 + j)
    V[f"{c}{e0 + 1}"] = f"=({VSUM['Market ex-cash, ex-interest NTM P/E on S&P FY27 $1.62']}*{c}{VR['FY28 ex-interest EPS']}+{c}{VR['Cash per share at horizon']})/{PRICE}-1"
    V[f"{c}{e0 + 1}"].number_format = PCT
V.cell(e0 + 1, 1, "Return (EPS-only)")
V.cell(e0 + 2, 1, "Expected return for a SHORT (EPS-only)").font = BOLD
V[f"B{e0 + 2}"] = f"=-SUMPRODUCT(B{VR['Probability']}:E{VR['Probability']},B{e0 + 1}:E{e0 + 1})"
V[f"B{e0 + 2}"].number_format = PCT
V.column_dimensions["A"].width = 58
for c in "BCDE":
    V.column_dimensions[c].width = 22

# ---------------------------------------------------------------- Consensus
C = wb.create_sheet("Consensus")
C["A1"] = "Consensus reconciliation (secondary sources; pulled 2026-10-01) vs our base (W2)"
C["A1"].font = Font(bold=True, size=13)
header(C, 3, ["Provider", "Q1-FY27 EPS", "FY27 EPS", "FY28 EPS", "n (FY27)", "Updated", "ACV in estimates?", "Note"])
cons = [("Yahoo Finance", None, 1.75, None, 8, "partial post-print", "mixed", "Stale-mixed: $1.85 30d ago (pre-print); Q1 row is stale Q4-FY26"),
        ("S&P Global (StockAnalysis)", None, 1.62, None, 9, "2026-09-25", "No (FY27 revenue $4.8bn)", "Governing FY27 basis"),
        ("Zacks (Nasdaq)", 0.41, 1.65, 1.81, 5, "2026-10-01", "Unclear", "Governing Q1 basis (n=4, $0.38-0.43)")]
for i, row in enumerate(cons, start=4):
    for j, v in enumerate(row):
        cc = C.cell(i, 1 + j, v)
        if isinstance(v, float):
            cc.number_format = NUM2
            cc.font = INPUT
C.cell(8, 1, "Our base (W2)").font = BOLD
C["B8"] = "=EPS_Build!B" + str(3 + 1 * 18 + 2 + 12)   # W2 Q1 EPS (no ACV in Q1)
C["C8"] = "=" + EPSX_CELLS[1]
C["D8"] = "=Valuation!C7"
for ref in ("B8", "C8", "D8"):
    C[ref].number_format = NUM2
C.cell(9, 1, "Gap vs governing consensus").font = BOLD
C["B9"] = "=B8/B6-1"   # vs Zacks Q1
C["C9"] = "=C8/C5-1"   # vs S&P FY27 (ex-ACV)
C["D9"] = "=D8/D6-1"   # vs Zacks FY28
for ref in ("B9", "C9", "D9"):
    C[ref].number_format = PCT
C.cell(10, 1, "Note: Q1 has no ACV P&L (close ~mid-Oct). FY27 gap uses our ex-ACV EPS to match S&P's ex-ACV basis. FY28 compares our GAAP incl. ACV with Zacks.")
C.column_dimensions["A"].width = 30
for c in "BCDEFGH":
    C.column_dimensions[c].width = 18
C.column_dimensions["H"].width = 60

# ---------------------------------------------------------------- Sensitivity
T = wb.create_sheet("Sensitivity")
T["A1"] = "Sensitivity of SHORT expected return to cost of equity and probabilities (live)"
T["A1"].font = Font(bold=True, size=13)
header(T, 3, ["Cost of equity shift", "Short expected return"])
for k, sh in enumerate([-0.01, -0.005, 0.0, 0.005, 0.01]):
    rr = 4 + k
    T.cell(rr, 1, sh).number_format = "+0.0%;-0.0%;0.0%"
    terms = []
    for j in range(4):
        c = get_column_letter(2 + j)
        g, ro, gt = f"Valuation!{c}{VR['5-yr growth g']}", f"Valuation!{c}{VR['Incremental ROIC']}", GT
        r = f"({COE}+A{rr})"
        pv5 = "+".join(f"(1+{g})^{t - 1}*(1-{g}/{ro})/(1+{r})^{t}" for t in range(1, 6))
        tv = f"(1+{g})^5*(1-{gt}/{ro})/({r}-{gt})/(1+{r})^5"
        val = f"(({pv5}+{tv})*Valuation!{c}{VR['FY28 ex-interest EPS']}+Valuation!{c}{VR['Cash per share at horizon']})"
        terms.append(f"Valuation!{c}{VR['Probability']}*({val}/{PRICE}-1)")
    T[f"B{rr}"] = "=-(" + "+".join(terms) + ")"
    T[f"B{rr}"].number_format = PCT
T.column_dimensions["A"].width = 24
T.column_dimensions["B"].width = 24

# ---------------------------------------------------------------- Catalysts & Sources
K = wb.create_sheet("Catalysts")
header(K, 1, ["Date", "Event", "Known today", "New information", "Supports", "Move needed to matter"])
cat = [("2026-10-07 / 10-13", "ACV tender expiry / HSR waiting period (SC TO-T/A)", "Terms; 47.8% tendered; HSR refiled 9/28",
        "Second Request?", "Mildly long if delayed", "<2%"),
       ("est. early Nov", "RB Global Q3-26", "Partner in all 50 states", "IAA units, first full quarter", "Short if IAA strong", "Read-across"),
       ("2026-11-19 (Zacks est.)", "CPRT Q1-FY27", "Lost customer (qualitative)",
        "Full-quarter US insurance units; OpEx/car; operating margin", "Short if US ins. ≤ −10% & op. margin < 34%",
        "≥7-10% NTM EPS cut (80th-pct move 6.7%)"),
       ("est. mid-Feb 2027", "CPRT Q2-FY27 + ACV consolidation", "HSBC −3% dilution; mgmt 'breakeven'",
        "PPA amortization; ACV losses in GAAP EPS", "Short if dilution > 3%", "2-3% incremental EPS"),
       ("est. late May 2027", "CPRT Q3-FY27", "—", "Last pre-lap quarter", "Ex-customer trend decides", "—"),
       ("est. Sep 2027", "CPRT Q4-FY27", "Lap is mechanical", "Growth resumption", "Long", "Confirms consensus")]
for i, row in enumerate(cat, start=2):
    for j, v in enumerate(row):
        K.cell(i, 1 + j, v).alignment = Alignment(wrap_text=True, vertical="top")
for c, w in zip("ABCDEF", [20, 30, 30, 36, 30, 30]):
    K.column_dimensions[c].width = w
Z = wb.create_sheet("Sources")
header(Z, 1, ["Datapoint", "Value", "Source (primary unless marked)", "Date"])
src = [("FY26 segment P&L (US/Intl)", "see Inputs", "Copart 10-K FY26 segment note; 10-Qs Q1-Q3 FY26", "2026-09-29"),
       ("Q4-FY26 US insurance units", "−7.5% (+2.3% ex one customer)", "Copart Q4-FY26 earnings call", "2026-09-10"),
       ("OpEx per car", "+12.7% y/y", "Copart Q4-FY26 earnings call", "2026-09-10"),
       ("ACV deal", "$10.50/sh cash, ~$1.9bn", "Copart 8-K / SC TO-T", "2026-09-10/17"),
       ("ACV TTM financials", "Rev $801m; op. loss −$57m", "ACV Auctions 10-Q XBRL", "2026-Q2"),
       ("RBA partner expansion", "All 50 states, 90-day transition", "RB Global Q2-26 call", "2026-08-04"),
       ("Progressive policies in force", "+7-9% y/y", "Progressive Aug-26 monthly report", "2026-09-18"),
       ("Manheim index", "−0.4% y/y mid-Sep", "Cox Automotive (secondary)", "2026-09"),
       ("Consensus", "see Consensus tab", "S&P Global via StockAnalysis; Zacks via Nasdaq; Yahoo (secondary)", "2026-10-01"),
       ("10y UST", "5.24%", "^TNX (secondary)", "2026-10-01"),
       ("Interest income FY26", "$181.9m", "Copart 10-K FY26", "2026-09-29")]
for i, row in enumerate(src, start=2):
    for j, v in enumerate(row):
        Z.cell(i, 1 + j, v)
for c, w in zip("ABCD", [32, 30, 60, 14]):
    Z.column_dimensions[c].width = w

# ---------------------------------------------------------------- Summary (first)
M = wb.create_sheet("Summary", 0)
M["A1"] = "CPRT — SHORT — 3-12 month horizon (close after Q2-FY27, mid-Feb 2027)"
M["A1"].font = Font(bold=True, size=14)
M["A3"] = "Thesis"
M["B3"] = ("The Street knows Copart lost Progressive's salvage volume but has not put it in the numbers: consensus FY27 EPS "
           "still grows; we model the first full-run-rate quarters (Q1-Q2 FY27) taking EPS ~11% below consensus, with "
           "ACV adding GAAP dilution. At 5.2% rates the stock is fairly valued only if Street EPS is right.")
M["B3"].alignment = Alignment(wrap_text=True, vertical="top")
M.row_dimensions[3].height = 75
items = [("Share price", f"={PRICE}", USD), ("Base-case (W2) price target", f"={VSUM['Base-case (W2) price target']}", USD),
         ("Probability-weighted value", f"={VSUM['Probability-weighted value per share']}", USD),
         ("Expected return for the short", f"={VSUM['Expected return for a SHORT']}", PCT),
         ("Expected return for the short with NO multiple change", f"=Valuation!B{e0 + 2}", PCT),
         ("Maximum plausible loss (worst world)", f"={VSUM['Maximum plausible loss for a short (worst world)']}", PCT),
         ("Our FY27 EPS ex-ACV vs S&P consensus $1.62", "=Consensus!C9", PCT),
         ("Our Q1-FY27 EPS vs Zacks consensus $0.41", "=Consensus!B9", PCT),
         ("Falsification: exit if Q1-FY27 US insurance units better than", -0.06, PCT)]
for k, (lab, formula, f) in enumerate(items, start=5):
    M.cell(k, 1, lab)
    M[f"B{k}"] = formula
    M[f"B{k}"].number_format = f
    M[f"B{k}"].font = BOLD if k < 13 else INPUT
M.column_dimensions["A"].width = 58
M.column_dimensions["B"].width = 95

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
out = ROOT / "model" / "TeamName.xlsx"
wb.save(out)
print("saved", out)
