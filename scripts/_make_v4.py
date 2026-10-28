"""Create build_final_model_v4.py from build_final_model_v3.py (v4: year-specific payout, full-sample event study,
source fixes, W2 waterfall + CoE chart). v2/v3 scripts and files untouched."""
from pathlib import Path

here = Path(__file__).parent
s = (here / "build_final_model_v3.py").read_text(encoding="utf-8")


def rep(a, b, count=1):
    global s
    assert s.count(a) >= 1, a[:90]
    s = s.replace(a, b) if count == 0 else s.replace(a, b, count)


rep('"""Build final/Margin_Callers_v3.xlsx (v3 = surgical QA fixes on v2)', '"""Build final/Margin_Callers_v4.xlsx (v4 = year-specific payout, full-sample event study, charts)')
rep('out = ROOT / "final" / "Margin_Callers_v3.xlsx"', 'out = ROOT / "final" / "Margin_Callers_v4.xlsx"')

# ---------------------------------------------------------------- 1. year-specific payout
rep('''def vexc_f(e1, e2, g, ro, r_, gt):
    """Ex-cash value: payout (1-g/ROIC) on year-1 (NTM) and year-2 earnings, years 3-5 at g, Gordon terminal."""
    po = f"(1-{g}/{ro})"
    pv = (f"{po}*({e1}/(1+{r_})+{e2}/(1+{r_})^2+{e2}*(1+{g})/(1+{r_})^3+{e2}*(1+{g})^2/(1+{r_})^4"
          f"+{e2}*(1+{g})^3/(1+{r_})^5)")''',
    '''def vexc_f(e1, e2, g, ro, r_, gt, po1=None):
    """Ex-cash value: year-1 (NTM) distribution at payout 1 - g1/ROIC (g1 = modeled year-2 / year-1 growth); years 2-5 at the
    normalized payout 1 - g/ROIC with years 3-5 growing at g; Gordon terminal."""
    po = f"(1-{g}/{ro})"
    p1 = po1 or f"(1-({e2}/{e1}-1)/{ro})"
    pv = (f"{p1}*{e1}/(1+{r_})+{po}*({e2}/(1+{r_})^2+{e2}*(1+{g})/(1+{r_})^3+{e2}*(1+{g})^2/(1+{r_})^4"
          f"+{e2}*(1+{g})^3/(1+{r_})^5)")''')
rep('("gt", "Terminal growth (Inputs)", PCT), ("po", "Payout = 1 − g / ROIC", PCT),',
    '("gt", "Terminal growth (Inputs)", PCT), ("po1", "Year-1 payout = 1 − g1 / ROIC (g1 = year-2 growth above)", PCT), '
    '("po", "Normalized payout, years 2-5 = 1 − g / ROIC", PCT),')
rep('pv, tv = vexc_f(ref("ntm"), ref("ntm2"), ref("g"), ref("roic"), ref("r"), ref("gt"))',
    'pv, tv = vexc_f(ref("ntm"), ref("ntm2"), ref("g"), ref("roic"), ref("r"), ref("gt"), po1=ref("po1"))')
rep('''"po": f"=1-{ref('g')}/{ref('roic')}",''', '''"po1": f"=1-{ref('y2g')}/{ref('roic')}", "po": f"=1-{ref('g')}/{ref('roic')}",''')
rep('"Value = PV of distributions (payout 1 − g/ROIC) from year-1 (NTM) and year-2 earnings, years 3-5 at g and a Gordon terminal, "',
    '"Value = PV of distributions from year-1 (NTM) and year-2 earnings, years 3-5 at g and a Gordon terminal, "')
rep('''"(derived from growth, ROIC and cost of equity, not selected).")''',
    '''"(derived from growth, ROIC and cost of equity, not selected). Explicit-period payout ratios are linked to each year's modeled EPS growth and "
      "incremental ROIC; normalized payout is used only after the recovery period.")''')

# ---------------------------------------------------------------- 4. Copart SEC-filed transcript; 6. Zacks dates; wording
TR = "https://www.sec.gov/Archives/edgar/data/900075/000119312526388383/d138398dex992.htm"
rep('("ins_rep", "Q4 FY26 domestic insurance assignments y/y (reported)", -0.075, PCT, "%", "Secondary transcript",',
    '("ins_rep", "Q4 FY26 domestic insurance assignments y/y (reported)", -0.075, PCT, "%", "Primary (SEC-filed transcript)",')
rep('("ins_excust", "Q4 FY26 domestic insurance assignments y/y, excluding one lost customer", 0.023, PCT, "%", "Secondary transcript",',
    '("ins_excust", "Q4 FY26 domestic insurance assignments y/y, excluding one lost customer", 0.023, PCT, "%", "Primary (SEC-filed transcript)",')
rep('"Q4 FY26 call (9/10/26): domestic insurance assignments −7.5%; CFO: U.S. insurance volumes −7.5%"',
    '"Q4 FY26 call (9/10/26; transcript filed as Schedule TO-C Ex. 99.2): domestic insurance assignments −7.5%; CFO: U.S. insurance volumes −7.5%"')
rep("OpEx/car stays high (+12.7% in Q4, transcript summary).", "OpEx/car stays high (+12.7% in Q4, call).")
rep("(Q4-FY26 +12.7%, transcript summary; U.S. facility cost per unit +14.2%, CFO)", "(Q4-FY26 +12.7%, CEO; U.S. facility cost per unit +14.2%, CFO; Q4 call)")
for q in ("Q1 FY27 (Oct-26)", "Q2 FY27 (Jan-27)", "Q3 FY27 (Apr-27)", "Q4 FY27 (Jul-27)", "FY27 (Jul-27)", "FY28 (Jul-28)"):
    rep(f'("Zacks: {q}", "10/2/26",', f'("Zacks: {q}", "10/1/26",')
rep('"Via Nasdaq (Zacks); quarterly consensus; unchanged vs 10/1/26 pull"', '"Via Nasdaq. Zacks estimates as of 10/1/26; re-verified 10/2/26."')
rep('header(C, 4, ["Provider / period", "Date pulled", "EPS ($/sh)", "n", "Low", "High", "Basis / note"])',
    'header(C, 4, ["Provider / period", "As of", "EPS ($/sh)", "n", "Low", "High", "Basis / note"])')
rep('''    ("Q4 domestic insurance assignments; ex-one-customer", "−7.5%; +2.3%", "Copart Q4 FY26 earnings call (transcript)", "Secondary transcript", "9/10/26",
     "The Motley Fool transcript (fool.com/earnings/call-transcripts); official webcast replay: copart.com/investorrelations", "Basis of the 9.8-pt drag (Inputs)"),
    ("OpEx per car; U.S. facility costs; total-loss frequency", "+12.7%; +7.7% (+14.2% per unit); 23.3% (record second quarter, CY Q2-26)",
     "Copart Q4 FY26 earnings call (transcript)", "Secondary transcript", "9/10/26",
     "The Motley Fool transcript; official webcast replay: copart.com/investorrelations", "+12.7% is from the transcript's takeaway summary; +14.2% is the CFO's per-unit figure"),''',
    f'''    ("Q4 domestic insurance assignments; ex-one-customer", "−7.5%; +2.3%", "Copart Q4 FY26 earnings-call transcript (Schedule TO-C, Ex. 99.2)",
     "Primary — SEC-filed earnings-call transcript, Ex. 99.2", "9/10/26 call; filed 9/11/26",
     "SC TO-C, accession 0001193125-26-388383: {TR}", "Basis of the 9.8-pt drag (Inputs)"),
    ("OpEx per car; U.S. facility costs; total-loss frequency", "+12.7% (CEO); +7.7% (+14.2% per unit, CFO); 23.3%, highest second quarter on record (CY Q2-26)",
     "Copart Q4 FY26 earnings-call transcript (Schedule TO-C, Ex. 99.2)", "Primary — SEC-filed earnings-call transcript, Ex. 99.2", "9/10/26 call; filed 9/11/26",
     "SC TO-C, accession 0001193125-26-388383", ""),''')
rep('''"Copart FY26 Form 10-K (Note 19; MD&A); Q4 FY26 call", "Primary", "9/29/26; 9/10/26", "10-K, accession 0001193125-26-405731; call per row 10", ""),''',
    '''"Copart FY26 Form 10-K (Note 19; MD&A); Q4 FY26 call transcript", "Primary", "9/29/26; 9/10/26",
     "10-K, accession 0001193125-26-405731; SC TO-C accession 0001193125-26-388383 (Ex. 99.2)", "Close date is a team assumption (mid-Q2)"),''')
rep('''    ("ACV EPS guidance (company)", "Neutral on EPS in the first full year; accretive FY28 and beyond", "Copart Form 8-K Ex. 99.1 (deal release)", "Primary", "9/10/26",
     "8-K, accession 0001193125-26-388064", "FY27 GAAP dilution and FY28 contribution are team estimates"),''',
    '''    ("ACV EPS guidance (company; inconsistent across materials)",
     "Formal transaction announcement: neutral to EPS in the first full year of ownership and accretive in FY2028 and beyond. Subsequent Q4 earnings-call "
     "language indicated first-full-year accretion. Model conservatively governs off the formal transaction announcement.",
     "Copart 8-K Ex. 99.1 / 99.2 (deal announcement); Q4 FY26 call transcript (SC TO-C Ex. 99.2)", "Primary", "9/10/26; 9/11/26",
     "8-K accession 0001193125-26-388064; SC TO-C accession 0001193125-26-388383", "Modeled FY27 ACV impact and FY28 contribution are team estimates"),
    ("ACV acquired net cash", "$242.3m cash − $205.0m line of credit = +$37.3m (6/30/26)", "ACV Auctions Form 10-Q (Q2-26), XBRL", "Primary", "8/10/26",
     "ACV 10-Q accession 0001637873-26-000034",
     "ACV acquired net cash is not explicitly modeled in horizon cash; June-30 ACV net cash was approximately +$37m before subsequent changes, an impact of < $0.05/share."),''')
rep('''"Primary (vehicles: Secondary transcript)", "9/29/26; 9/10/26",
     "10-K accession 0001193125-26-405731; 8-K accession 0001193125-26-387902",''',
    '''"Primary", "9/29/26; 9/10/26",
     "10-K accession 0001193125-26-405731; 8-K accession 0001193125-26-387902; SC TO-C accession 0001193125-26-388383",''')
rep('"Copart FY26 Form 10-K (Item 2 Properties); 8-K Ex. 99.1 (About Copart); Q4 FY26 call"',
    '"Copart FY26 Form 10-K (Item 2 Properties); 8-K Ex. 99.1 (About Copart); Q4 FY26 call transcript (SC TO-C Ex. 99.2)"')
rep('''("Consensus EPS (quarterly, FY27, FY28)", "See Consensus sheet", "Zacks via Nasdaq", "Secondary consensus", "10/2/26",''',
    '''("Consensus EPS (quarterly, FY27, FY28)", "See Consensus sheet", "Zacks via Nasdaq", "Secondary consensus", "10/1/26; re-verified 10/2/26",''')
rep('"Unchanged vs 10/1/26"', '"Zacks estimates as of 10/1/26; re-verified 10/2/26."')
rep('("Post-earnings-reaction study", "3M excess +4.6% / +6.0% / +6.9% (n = 21 / 12 / 8)"',
    '("Post-earnings-reaction study", "Full sample (governing): 3M mean excess +4.5% / +5.6% / +6.3% (n = 22 / 13 / 9)"')

# ---------------------------------------------------------------- 5. ACV net cash note on Valuation
rep('''"Horizon cash is held constant across scenarios to isolate operating differences; ±$200m of cash changes value by approximately $0.21/share.", font=F_NOTE)''',
    '''"Horizon cash is held constant across scenarios to isolate operating differences; ±$200m of cash changes value by approximately $0.21/share.", font=F_NOTE)
put(V, f"A{re_ + 7}", "ACV acquired net cash is not explicitly modeled in horizon cash; June-30 ACV net cash was approximately +$37m before subsequent changes, an impact of < $0.05/share.", font=F_NOTE)''')

# ---------------------------------------------------------------- 2. event study: full sample governs, ex-COVID robustness
a = s.index('section(X, ev, "Post-earnings-reaction study (holding-period discipline)", 7)')
b = s.index('put(X, f"A{er + 4}", "6M mean excess')
b = s.index("\n", b) + 1
s = s[:a] + '''section(X, ev, "Post-earnings-reaction study (holding-period discipline)", 7)
EVD = [
    ("Sample period", "Governing: all CPRT earnings releases Sep 2004 - Sep 2026 (89 releases), COVID period included. Robustness: excludes Feb 2020 - Jun 2023. The 9/10/26 release has no forward window yet."),
    ("Event definition", "2-day excess reaction vs SPY, close(t+1) / close(t−1), around the earnings-release 8-K date; a price reaction, not an EPS miss vs consensus"),
    ("Thresholds tested", "Excess reaction ≤ −3%, ≤ −5%, ≤ −7%"),
    ("Benchmark / forward window", "SPY; excess return from the close of day t+1 over 21 / 63 / 126 trading days (1M / 3M / 6M); 3M governs"),
    ("Source", "Team backtest (scripts/backtest.py); prices: Yahoo Finance; dates: SEC 8-K filing dates"),
]
for i, (a, b) in enumerate(EVD, start=ev + 1):
    put(X, f"A{i}", a)
    X.merge_cells(f"B{i}:G{i}")
    put(X, f"B{i}", b, wrap=True)
er = ev + len(EVD) + 2
EVT = [("Governing: full sample (incl. COVID)", [("≤ −3%", 23, 22, -0.0062, 0.0448, 0.0318, 0.7273), ("≤ −5%", 14, 13, -0.0134, 0.0564, 0.0628, 0.7692),
                                                ("≤ −7%", 9, 9, -0.0225, 0.0629, 0.0308, 0.7778)]),
       ("Robustness only: ex-COVID (Feb 2020 - Jun 2023 excluded)", [("≤ −3%", 22, 21, -0.0038, 0.0461, 0.0328, 0.7143), ("≤ −5%", 13, 12, -0.0099, 0.0596, 0.0667, 0.75),
                                                                     ("≤ −7%", 8, 8, -0.0183, 0.0685, 0.0525, 0.75)])]
for lab_t, rows_t in EVT:
    put(X, f"A{er}", lab_t, font=F_BOLD)
    header(X, er + 1, ["Threshold (2-day excess)", "Events", "n (with 3M)", "1M mean", "3M mean", "3M median", "3M hit rate"])
    for i, row_ in enumerate(rows_t, start=er + 2):
        put(X, f"A{i}", row_[0])
        put(X, f"B{i}", row_[1], "0", F_IN, align="right")
        put(X, f"C{i}", row_[2], "0", F_IN, align="right")
        for k, v in enumerate(row_[3:]):
            put(X, f"{get_column_letter(4 + k)}{i}", v, PCT, F_IN, align="right")
    er += 6
put(X, f"A{er - 1}", "Full-sample 6M mean excess: +3.6% / +4.3% / +4.1%. Positive 3M drift after large negative reactions in both samples argues against holding the short past Q2.", font=F_NOTE)
''' + s[b:]

# ---------------------------------------------------------------- payout validity checks
rep('''    ("ACV consolidated share of Q2 between 0% and 100%", f'=IF(AND({NM["acv_frac"]}>=0,{NM["acv_frac"]}<=1),"OK","CHECK")'),''',
    '''    ("ACV consolidated share of Q2 between 0% and 100%", f'=IF(AND({NM["acv_frac"]}>=0,{NM["acv_frac"]}<=1),"OK","CHECK")'),
    ("Payout ratios valid: year-1 and normalized payouts between 0% and 100%, all worlds",
     f'=IF(AND(MIN(Valuation!B{VR["po1"]}:E{VR["po1"]},Valuation!B{VR["po"]}:E{VR["po"]})>0,MAX(Valuation!B{VR["po1"]}:E{VR["po1"]},Valuation!B{VR["po"]}:E{VR["po"]})<=1),"OK","CHECK")'),
    ("Growth below incremental ROIC (year-1 g1 and normalized g), all worlds",
     f'=IF(AND(MAX(Valuation!B{VR["y2g"]}-Valuation!B{VR["roic"]},Valuation!C{VR["y2g"]}-Valuation!C{VR["roic"]},Valuation!D{VR["y2g"]}-Valuation!D{VR["roic"]},Valuation!E{VR["y2g"]}-Valuation!E{VR["roic"]})<0,'
     f'MAX(Valuation!B{VR["g"]}-Valuation!B{VR["roic"]},Valuation!C{VR["g"]}-Valuation!C{VR["roic"]},Valuation!D{VR["g"]}-Valuation!D{VR["roic"]},Valuation!E{VR["g"]}-Valuation!E{VR["roic"]})<0),"OK","WARNING")'),
    ("Waterfall ties: FY26 EPS + bridge = FY27 ex-ACV; + ACV impact = FY27 GAAP (W2)",
     f'=IF(AND(ABS({TS["wf_exacv"]}-{EPSX[1]})<0.000001,ABS({TS["wf_gaap"]}-{EPSG[1]})<0.000001),"OK","CHECK")'),''')

# ---------------------------------------------------------------- 9. bridge moved to EPS_Build + waterfall chart data
a = s.index("# ---- W2 FY27 EPS bridge from FY26 reported EPS")
b = s.index("# ================================================================== Consensus")
s = s[:a] + r'''# ---- W2 FY27 EPS bridge (EPS_Build, below the W4 block): cumulative replay + waterfall chart data (all linked)
row = 4 + 4 * BLOCK + 1
section(E, row, "W2 FY27 EPS bridge from FY26 reported EPS (cumulative replay of the W2 block; drivers switched on in order)", 11)
header(E, row + 1, ["Step", "Underlying u", "Run-rate drag", "Q4 base drag on", "RPU growth", "Fixed cost growth", "Non-ins. growth",
                    "Veh. GP / G&A growth", "Intl growth", "FY27 EPS ex-ACV", "Step change"])
BR0 = row + 2
STEPS = [
    ("FY26 reported diluted EPS", None),
    ("FY26 replay at FY27 share count and pre-ACV interest (zero growth)", dict()),
    ("+ lost-customer drag (13-pt run-rate, net of Q4 FY26 base)", dict(L=1, D=1)),
    ("+ underlying insurance and non-insurance unit growth", dict(L=1, D=1, U=1, G=1)),
    ("+ revenue per unit", dict(L=1, D=1, U=1, G=1, R=1)),
    ("+ fixed facility cost growth", dict(L=1, D=1, U=1, G=1, R=1, F=1)),
    ("+ international, vehicle gross profit and G&A growth (= W2)", dict(L=1, D=1, U=1, G=1, R=1, F=1, H=1, I=1)),
]
SRCREF = {"U": sc2("u"), "L": sc2("loss"), "D": "1", "R": sc2("rpu"), "F": sc2("fixed"), "G": NM["nonins"], "H": NM["goth"], "I": NM["gint"]}
for i, (lab, on) in enumerate(STEPS):
    rr = BR0 + i
    put(E, f"A{rr}", lab)
    if on is None:
        put(E, f"J{rr}", f"=SUM(Inputs!B{SEG['eps26']}:E{SEG['eps26']})", USD, F_LINK, align="right")
        continue
    for col, key in zip("BCDEFGHI", "ULDRFGHI"):
        if key in on:
            put(E, f"{col}{rr}", f"={SRCREF[key]}" if key != "D" else 1, PCT if key != "D" else "0", F_LINK if key != "D" else F_IN, align="right")
        else:
            put(E, f"{col}{rr}", 0, PCT if key != "D" else "0", F_IN, align="right")
    terms = []
    for q in range(4):
        c = get_column_letter(2 + q)
        g = lambda k: f"Inputs!${c}${SEG[k]}"
        imp = f"((1+$C{rr})/(1-{g('lib')}*$D{rr})-1)"
        units = f"({NM['sins']}*($B{rr}+{imp})+(1-{NM['sins']})*$G{rr})"
        usoi = (f"({g('us_service')}*(1+{units})*(1+$E{rr})+({g('us_vehicle')}-{g('us_cogs_vehicle')})*(1+$H{rr})"
                f"-({g('us_facility')}*{NM['fix']}*(1+$F{rr})+{g('us_facility')}*(1-{NM['fix']})*(1+{units}))-{g('us_ga')}*(1+$H{rr}))")
        terms.append(f"({usoi}+{g('intl_oi')}*(1+$I{rr}))")
    put(E, f"J{rr}", "=(" + "+".join(terms) + f"+4*{NM['intq1']})*(1-{NM['tax']})/{NM['shares']}", USD, F_BOLD if i == 6 else F_BODY, align="right")
    put(E, f"K{rr}", f"=J{rr}-J{rr - 1}", USD, align="right")
TS["bridge_end"] = f"EPS_Build!$J${BR0 + 6}"
row = BR0 + len(STEPS) + 1
section(E, row, "Chart data: W2 EPS waterfall (linked; memo bridge = column B at $0.01)", 11)
header(E, row + 1, ["Bridge line", "$/sh", "Level after", "Base (hidden)", "Total", "Increase", "Decrease"])
WF0 = row + 2
WF = [("FY26 EPS", "t", f"=J{BR0}"), ("Shares / interest / other", "d", f"=K{BR0 + 1}+K{BR0 + 6}"),
      ("Lost customer", "d", f"=K{BR0 + 2}"), ("Underlying growth", "d", f"=K{BR0 + 3}"), ("Revenue per unit", "d", f"=K{BR0 + 4}"),
      ("Fixed-cost deleverage", "d", f"=K{BR0 + 5}"), ("FY27 ex-ACV", "t", f"={EPSX[1]}"),
      ("ACV impact", "d", f"=EPS_Build!$F${QX_ROWS[1]['acv_eps']}"), ("FY27 GAAP", "t", f"={EPSG[1]}")]
for i, (lab, kind, f) in enumerate(WF):
    rr = WF0 + i
    put(E, f"A{rr}", lab, font=F_BOLD if kind == "t" else F_BODY)
    put(E, f"B{rr}", f, USD, F_BOLD if kind == "t" else F_BODY, align="right")
    put(E, f"C{rr}", f"=B{rr}" if kind == "t" else f"=C{rr - 1}+B{rr}", USD, align="right")
    put(E, f"D{rr}", 0 if kind == "t" else f"=MIN(C{rr - 1},C{rr})", USD, F_IN if kind == "t" else F_BODY, align="right")
    put(E, f"E{rr}", f"=B{rr}" if kind == "t" else 0, '$0.00;;;', F_BODY if kind == "t" else F_IN, align="right")
    put(E, f"F{rr}", 0 if kind == "t" else f"=MAX(B{rr},0)", '"+"$0.00;;;', F_IN if kind == "t" else F_BODY, align="right")
    put(E, f"G{rr}", 0 if kind == "t" else f"=MAX(-B{rr},0)", '"−"$0.00;;;', F_IN if kind == "t" else F_BODY, align="right")
TS["wf_exacv"] = f"EPS_Build!$C${WF0 + 5}"
TS["wf_gaap"] = f"EPS_Build!$C${WF0 + 7}"
TS["bridge_disp_ok"] = "+".join(f"ROUND(EPS_Build!$B${WF0 + k},2)" for k in range(6)) + f"-ROUND(EPS_Build!$B${WF0 + 6},2)"

''' + s[b:]

# ---------------------------------------------------------------- CoE rows reference for the line chart
rep('''for k, shf in enumerate([-0.01, -0.005, 0.0, 0.005, 0.01]):
    rr = row + k''', '''COE0 = row
for k, shf in enumerate([-0.01, -0.005, 0.0, 0.005, 0.01]):
    rr = row + k''')

# ---------------------------------------------------------------- Summary notes / chart 1 series name
rep('put(M, "A25", "Value at mid-Feb 2027 on NTM (Q3 FY27-Q2 FY28) ex-interest EPS; P/E derived from growth, ROIC and cost of equity, not selected. Team scenario probabilities.", font=F_NOTE)',
    'put(M, "A25", "Value at mid-Feb 2027 on NTM (Q3 FY27-Q2 FY28) ex-interest EPS; P/E derived from growth, ROIC and cost of equity (year-specific payouts), not selected. Team scenario probabilities.", font=F_NOTE)')
rep('for c, lab in zip("ABC", ["Quarter", "Zacks consensus", "Our ex-ACV estimate"]):', 'for c, lab in zip("ABC", ["Quarter", "Zacks consensus", "Margin Callers ex-ACV"]):')

# ---------------------------------------------------------------- 8. Summary chart 2 rebuilt with a linked price reference line
a = s.index("# Chart 2: horizontal bars of scenario value")
b = s.index('M.column_dimensions["I"].width = 3')
s = s[:a] + r'''# Chart 2: horizontal bars of scenario value + linked current-price reference line (scatter on hidden secondary axes)
from openpyxl.chart import ScatterChart, LineChart, Series
from openpyxl.chart.series import SeriesLabel, StrRef
from openpyxl.chart.layout import Layout, ManualLayout
put(M, "A41", "Chart data: current-price reference line (linked to Inputs)", font=F_SEC)
for c in range(1, 4):
    M.cell(41, c).border = SEC_LINE
for c, lab in zip("ABC", ["Series label", "Price ($/sh)", "Height"]):
    put(M, f"{c}42", lab, font=F_HDR, border=HDR_LINE, align="left" if c == "A" else "right")
put(M, "A43", f'="Current price "&TEXT({NM["price"]},"$0.00")', border=ROW_LINE)
put(M, "A44", "", border=ROW_LINE)
for rr, h in ((43, 0), (44, 1)):
    put(M, f"B{rr}", f"={NM['price']}", USD, F_LINK, align="right", border=ROW_LINE)
    put(M, f"C{rr}", h, "0", F_IN, align="right", border=ROW_LINE)
ch2 = BarChart()
ch2.type = "bar"
ch2.grouping = "clustered"
ch2.title = _title("Scenario Value per Share at Mid-Feb 2027 Exit")
ch2.add_data(Reference(M, min_col=7, min_row=20, max_row=23), titles_from_data=False)
ch2.set_categories(Reference(M, min_col=1, min_row=20, max_row=23))
_clean(ch2)
ch2.legend = None
ch2.x_axis.scaling.orientation = "maxMin"          # W1 at top
ch2.y_axis.crosses = "max"                          # value axis at bottom when categories are reversed
ch2.y_axis.scaling.min = 0
ch2.y_axis.scaling.max = 35
ch2.y_axis.majorUnit = 5
ch2.y_axis.numFmt = "$0"
ch2.y_axis.majorGridlines = None
ch2.y_axis.title = _title("Value per share at mid-Feb 2027 ($)", size=800)
s2 = ch2.series[0]
s2.graphicalProperties = GraphicalProperties(solidFill="17365D", ln=LineProperties(noFill=True))
for idx, col in enumerate(["17365D", "17365D", "8C96A5", "8C96A5"]):
    pt = DataPoint(idx=idx)
    pt.graphicalProperties = GraphicalProperties(solidFill=col, ln=LineProperties(noFill=True))
    s2.dPt.append(pt)
s2.dLbls = DataLabelList()
s2.dLbls.showVal = True
s2.dLbls.showSerName = s2.dLbls.showCatName = s2.dLbls.showLegendKey = s2.dLbls.showPercent = False
s2.dLbls.position = "outEnd"
s2.dLbls.numFmt = "$0.00"
s2.dLbls.txPr = _rich(size=800, bold=True)
ch2.gapWidth = 60
pref = ScatterChart()
pref.scatterStyle = "lineMarker"
ps = Series(Reference(M, min_col=3, min_row=43, max_row=44), Reference(M, min_col=2, min_row=43, max_row=44))
ps.tx = SeriesLabel(strRef=StrRef("'Summary'!$A$43"))
ps.marker.symbol = "none"
ps.smooth = False
ps.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="9B2C2C", w=15875, prstDash="dash"))
from openpyxl.chart.label import DataLabel
ps.dLbls = DataLabelList()
ps.dLbls.showSerName = ps.dLbls.showVal = ps.dLbls.showCatName = ps.dLbls.showLegendKey = ps.dLbls.showPercent = False
top_lbl = DataLabel(idx=1, dLblPos="r", showSerName=True, showVal=False, showCatName=False, showLegendKey=False, showPercent=False)
top_lbl.txPr = _rich(size=800, bold=True, color="9B2C2C")
ps.dLbls.dLbl = [top_lbl]
pref.series.append(ps)
pref.x_axis.axId = 300
pref.y_axis.axId = 200
pref.x_axis.crossAx = 200
pref.y_axis.crossAx = 300
pref.x_axis.scaling.min, pref.x_axis.scaling.max = 0, 35
pref.y_axis.scaling.min, pref.y_axis.scaling.max = 0, 1
pref.x_axis.axPos, pref.y_axis.axPos = "t", "r"
pref.x_axis.delete = True
pref.y_axis.delete = True
pref.x_axis.majorGridlines = None
pref.y_axis.majorGridlines = None
pref.y_axis.crosses = "max"
ch2 += pref
ch2.height, ch2.width = 7.4, 15.0
M.add_chart(ch2, "J21")
''' + s[b:]

# ---------------------------------------------------------------- 9 + 10. waterfall (EPS_Build) and CoE line chart (Sensitivity)
rep('# ================================================================== global visual pass (styling only)', r'''# W2 EPS bridge waterfall (EPS_Build, right of the W2 block); every series links to the chart-data table
wf = BarChart()
wf.type = "col"
wf.grouping = "stacked"
wf.overlap = 100
wf.gapWidth = 55
wf.title = _title("W2 FY27 EPS Bridge: FY26 Actual to FY27E")
wf.add_data(Reference(E, min_col=4, max_col=7, min_row=WF0 - 1, max_row=WF0 + len(WF) - 1), titles_from_data=True)
wf.set_categories(Reference(E, min_col=1, min_row=WF0, max_row=WF0 + len(WF) - 1))
_clean(wf)
wf.legend = None
wf.y_axis.numFmt = "$0.00"
wf.y_axis.scaling.min = 1.2
wf.y_axis.scaling.max = 1.6
wf.y_axis.majorUnit = 0.1
wf.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E3E3E3", w=6350)))
wf.x_axis.txPr = _rich(size=700)
for sr, fill, fmt in zip(wf.series, [None, "17365D", "8C96A5", "9B2C2C"], [None, "$0.00;;;", '"+"$0.00;;;', '"-"$0.00;;;']):
    if fill is None:
        sr.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
        continue
    sr.graphicalProperties = GraphicalProperties(solidFill=fill, ln=LineProperties(noFill=True))
    sr.dLbls = DataLabelList()
    sr.dLbls.showVal = True
    sr.dLbls.showSerName = sr.dLbls.showCatName = sr.dLbls.showLegendKey = sr.dLbls.showPercent = False
    sr.dLbls.position = "inEnd" if fill == "17365D" else "ctr"
    sr.dLbls.numFmt = fmt
    sr.dLbls.txPr = _rich(size=750, bold=True, color="FFFFFF" if fill == "17365D" else "000000")
wf.height, wf.width = 8.6, 17.0
E.add_chart(wf, f"H{4 + BLOCK}")

# Cost-of-equity sensitivity line (Sensitivity, right of the table)
cl = ScatterChart()
cl.scatterStyle = "lineMarker"
cl.title = _title("Expected Short Return vs Cost of Equity")
cs = Series(Reference(T, min_col=5, min_row=COE0, max_row=COE0 + 4), Reference(T, min_col=3, min_row=COE0, max_row=COE0 + 4), title="Expected short return")
cs.marker.symbol = "circle"
cs.marker.size = 5
cs.marker.graphicalProperties = GraphicalProperties(solidFill="17365D", ln=LineProperties(solidFill="17365D"))
cs.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="17365D", w=19050))
cs.smooth = False
cs.dLbls = DataLabelList()
cs.dLbls.showVal = True
cs.dLbls.showSerName = cs.dLbls.showCatName = cs.dLbls.showLegendKey = cs.dLbls.showPercent = False
cs.dLbls.position = "t"
cs.dLbls.numFmt = '+0.0%;-0.0%;0.0%'
cs.dLbls.txPr = _rich(size=750)
cl.series.append(cs)
_clean(cl)
cl.legend = None
cl.x_axis.title = _title("Cost of equity", size=800)
cl.y_axis.title = _title("Expected short return", size=800)
cl.x_axis.numFmt = "0.0%"
cl.y_axis.numFmt = "0%"
cl.x_axis.scaling.min, cl.x_axis.scaling.max, cl.x_axis.majorUnit = 0.085, 0.11, 0.005
cl.y_axis.scaling.min, cl.y_axis.scaling.max, cl.y_axis.majorUnit = -0.10, 0.25, 0.05
cl.x_axis.tickLblPos = "low"
cl.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="595959", w=9525))
cl.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E3E3E3", w=6350)))
cl.height, cl.width = 7.0, 12.5
T.add_chart(cl, f"G{COE0 - 2}")

# ================================================================== global visual pass (styling only)''')

# print areas include the chart columns
rep("# Summary print area includes the charts (J:R)", '''E.print_area = f"A1:Q{E.max_row}"
# Summary print area includes the charts (J:R)''')
(here / "build_final_model_v4.py").write_text(s, encoding="utf-8")
print("ok")
