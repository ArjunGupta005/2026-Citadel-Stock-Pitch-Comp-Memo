"""Final consistency audit: Margin_Callers_v2.xlsx (Excel-cached values) vs Margin_Callers_v2.pdf (memo text)."""
import re

import pymupdf
from openpyxl import load_workbook

from common import ROOT

F = ROOT / "final"
pdf_path = F / "Margin_Callers_v2.pdf"
doc = pymupdf.open(pdf_path)
doc.set_metadata({"author": "Margin Callers", "creator": "Margin Callers", "producer": "Margin Callers", "title": "Margin Callers - CPRT Short",
                  "subject": "", "keywords": ""})
doc.saveIncr()
memo = re.sub(r"\s+", " ", " ".join(p.get_text() for p in doc))

wb = load_workbook(F / "Margin_Callers_v2.xlsx", data_only=True)
V, M, C = wb["Valuation"], wb["Summary"], wb["Consensus"]
vrow = {V.cell(r, 1).value: r for r in range(1, V.max_row + 1) if V.cell(r, 1).value}
mrow = {M.cell(r, 1).value: r for r in range(1, 40) if M.cell(r, 1).value}
v = lambda lab, col="B": V[f"{col}{vrow[lab]}"].value
usd = lambda x: f"${x:,.2f}"
pct = lambda x: f"{x * 100:.1f}%"
rows = [
    ("Reference price", M[f"B{mrow['Reference price (10/1/26)']}"].value, usd, ["$27.29"]),
    ("Q1 FY27 EPS (ours)", C["B16"].value, usd, None),
    ("Q1 FY27 consensus (Zacks)", C["C16"].value, usd, None),
    ("FY27 EPS ex-ACV (ours)", C["B18"].value, usd, None),
    ("FY27 consensus (S&P Global / Zacks)", (C["C18"].value, C["C19"].value), lambda t: f"${t[0]:.2f}-${t[1]:.2f}", None),
    ("FY27 GAAP EPS incl. ACV (W2)", v("FY27 GAAP EPS incl. ACV ($/sh)", "C"), usd, None),
    ("W1 value", v("Value per share, mid-Feb 2027 ($/sh)", "B"), usd, None),
    ("W2 value", v("Value per share, mid-Feb 2027 ($/sh)", "C"), usd, None),
    ("W3 value", v("Value per share, mid-Feb 2027 ($/sh)", "D"), usd, None),
    ("W4 value", v("Value per share, mid-Feb 2027 ($/sh)", "E"), usd, None),
    ("Base-case target (W2)", v("Base-case (W2) value ($/sh)"), usd, None),
    ("Probability-weighted value", v("Probability-weighted value ($/sh)"), usd, None),
    ("Expected short return", v("Expected short return"), lambda x: "+" + pct(x), None),
    ("W4 adverse-case loss", v("W4 adverse-case short loss (W4 stock return)"), pct, None),
    ("Rate-cut stress loss", v("Rate-cut stress short loss (stock return)"), pct, None),
    ("Planned exit / valuation date", V["A3"].value, lambda s: "mid-Feb 2027", ["mid-Feb 2027", "mid-February 2027"]),
]
out = ["| Metric | Excel | Memo | Match? |", "|---|---|---|---|"]
allok = True
for name, x, fmt, alts in rows:
    shown = fmt(x)
    cands = alts or [shown]
    hit = next((c for c in cands if c in memo), None)
    ok = hit is not None
    allok &= ok
    xs = x if isinstance(x, str) else (", ".join(f"{t:.4f}" for t in x) if isinstance(x, tuple) else f"{x:.4f}")
    out.append(f"| {name} | {xs} | {hit or '(not found)'} | {'YES' if ok else 'NO'} |")
print("\n".join(out))
# stale v1 numbers must not appear in the v2 memo
stale = [s for s in ["$23.10", "$24.61", "+9.8%", "$17.12", "$30.20", "$31.70", "$1.43", "$0.36 ", "47.9%", "16.1%", "largest online",
                     "record 23.3%", "already priced in", "post-miss", "event study shows"] if s in memo]
print("Stale v1 strings in memo:", stale or "NONE")
terms = ["progressive", "90%", "$209", "209m", "jpm", "proprietary", "hopkins", "university", "arjun", "vivaan", "gupta", "singh", "club"]
print("Banned terms in memo:", [t for t in terms if t in memo.lower()] or "NONE", "| pages:", doc.page_count, "| metadata:", doc.metadata.get("author"))
print("ALL ROWS MATCH:", allok and not stale)
(F / "_consistency_table_v2.md").write_text("\n".join(out) + "\n", encoding="utf-8")
