"""QA for final/Margin_Callers.xlsx: recalculation vs Python reference, error scan, fills, hidden sheets, banned terms."""
import re
import formulas
from openpyxl import load_workbook
import scenario_engine as se
import cprt_framework_b as fb
from common import ROOT

path = ROOT / "final" / "Margin_Callers.xlsx"
sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
vals = {}
for k, v in sol.items():
    m = re.search(r"\]([A-Z_]+)'!([A-Z]+\d+)$", k.upper())
    if m:
        try:
            vals[(m.group(1), m.group(2))] = v.value[0][0]
        except Exception:
            pass
errs = [(s, c, x) for (s, c), x in vals.items() if isinstance(x, str) and x.startswith("#")]
print("Formula error cells:", errs if errs else "NONE")
g = lambda s, c: float(vals[(s.upper(), c)])
W4 = dict(u=0.03, loss=-0.05, rpu=0.05, fixed=0.04, acv_amort=10, acv_op=0, int_q=28)
ok = True
for j, w in enumerate([se.CPRT_WORLDS["bear"], se.CPRT_WORLDS["base"], se.CPRT_WORLDS["bull"], W4]):
    col = "BCDE"[j]
    ref_g = sum(q["eps"] for q in se.cprt_quarters(w))
    ref_x = sum(q["eps"] for q in se.cprt_quarters(w, include_acv=False))
    xg, xx = g("Valuation", f"{col}7"), g("Valuation", f"{col}8")
    print(f"W{j+1}: GAAP {xg:.4f} vs py {ref_g:.4f} | ex-ACV {xx:.4f} vs py {ref_x:.4f} | value {g('Valuation', col+'20'):.2f} | ret {g('Valuation', col+'21'):+.4f}")
    ok &= abs(xg - ref_g) < 1e-9 and abs(xx - ref_x) < 1e-9
wb = load_workbook(path)
M = wb["Summary"]
labels = {M[f"A{r}"].value: r for r in range(1, 40) if M[f"A{r}"].value}
for lab in ["Reference price (10/1/26)", "Base-case value (W2)", "Probability-weighted value", "Expected short return",
            "Expected short return, EPS-only (no multiple change)", "W4 adverse-case short loss", "Rate-cut stress (10Y to 3.5%): short loss"]:
    print(f"Summary | {lab}: {g('Summary', 'B' + str(labels[lab])):.4f}")
for r in (12, 13, 14):
    print("Summary EPS row:", M[f"D{r}"].value, round(g("Summary", f"E{r}"), 4), round(g("Summary", f"F{r}"), 4), round(g("Summary", f"G{r}"), 4))
print("Checks:", vals.get(("CHECKS", "B11")), vals.get(("CHECKS", "B12")), "| calibration diff FY26:", round(g("Checks", "F7"), 4))
print("Sensitivity CoE shifts:", [round(g("Sensitivity", f"C{r}"), 4) for r in range(6, 11)],
      "| prob sets:", [round(g("Sensitivity", f"B{r}"), 4) for r in range(14, 17)])
print("ALL EPS MATCH PYTHON" if ok else "MISMATCH")
# formatting / hygiene
fills, comments = [], 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.fill is not None and c.fill.fill_type not in (None, "none"):
                fills.append((ws.title, c.coordinate))
            if c.comment:
                comments += 1
    assert ws.sheet_state == "visible", ws.title
    if ws.conditional_formatting:
        print("conditional formatting on", ws.title)
print("Filled cells:", fills if fills else "NONE", "| sheets:", wb.sheetnames, "| all visible | cell comments:", comments)
print("Gridlines hidden on all sheets:", all(not ws.sheet_view.showGridLines for ws in wb.worksheets))
print("Defined names:", list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else wb.defined_names)
text = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None).lower()
text += " " + " ".join(str(getattr(wb.properties, a) or "") for a in ("creator", "lastModifiedBy", "title")).lower()
for term in ["progressive", "90%", "209", "jpm", "jpmorgan", "proprietary", "g² theorem", "g2 theorem", "theorem", "johns hopkins", "jhu",
             "hopkins", "university", "arjun", "vivaan", "gupta", "singh", "club", "ex-acv basis"]:
    hits = [m.start() for m in re.finditer(re.escape(term), text)]
    print(f"term '{term}': {len(hits)}" + (f" -> {text[max(0, hits[0]-60):hits[0]+60]!r}" if hits else ""))
print("workbook props:", wb.properties.creator, "|", wb.properties.lastModifiedBy, "|", wb.properties.title)
