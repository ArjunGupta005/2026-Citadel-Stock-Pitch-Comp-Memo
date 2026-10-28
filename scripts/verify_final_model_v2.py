"""QA for final/Margin_Callers_v2.xlsx: full recalculation (formulas engine) vs the independent Python reference
(cprt_model_v2.py), error scan, integrity checks, fills, hidden sheets, defined names, banned terms."""
import re
import sys
import zipfile
from copy import deepcopy

import formulas
from openpyxl import load_workbook

import cprt_model_v2 as mv
from common import ROOT

path = ROOT / "final" / (sys.argv[1] if len(sys.argv) > 1 else "Margin_Callers_v2.xlsx")
sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
vals = {}
for k, v in sol.items():
    m = re.search(r"\]([A-Z_]+)'!([A-Z]+\d+)$", k.upper())
    if m:
        try:
            vals[(m.group(1), m.group(2))] = v.value[0][0]
        except Exception:
            pass
errs = [(s, c, x) for (s, c), x in vals.items() if isinstance(x, str) and x.startswith("#")]
print("Formula error cells:", errs if errs else "NONE")
g = lambda s, c: float(vals[(s.upper(), c)])
wb = load_workbook(path)
fails = []


def close(name, a, b, tol=1e-9):
    if abs(a - b) > tol:
        fails.append(f"{name}: excel {a:.6f} vs python {b:.6f}")


def rows(sheet):
    ws = wb[sheet]
    return {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}


ref = mv.run()
VRw = rows("Valuation")
pick = lambda lab: VRw[lab]
MAP = {"FY27 EPS ex-ACV ($/sh)": "fy27x", "FY27 modeled ACV impact ($/sh)": "acv27", "FY27 GAAP EPS incl. ACV ($/sh)": "fy27g",
       "FY28 standalone CPRT EPS before ACV ($/sh)": "fy28_sa", "FY28 EPS incl. ACV ($/sh)": "fy28",
       "NTM ex-interest EPS at mid-Feb 2027 (year 1)": "ntm_x", "Year 2 ex-interest EPS (Q3 FY28 - Q2 FY29)": "ntm2_x",
       "Justified ex-cash P/E on NTM EPS": "jpe", "Value per share, mid-Feb 2027 ($/sh)": "val", "Stock return vs reference price": "ret"}
for j, n in enumerate(["W1", "W2", "W3", "W4"]):
    col = "BCDE"[j]
    for lab, key in MAP.items():
        close(f"{n} {key}", g("Valuation", f"{col}{pick(lab)}"), ref["res"][n][key])
    r = ref["res"][n]
    print(f"{n}: Q1x {r['q']['eps_x'][0]:.3f} FY27x {r['fy27x']:.3f} ACV27 {r['acv27']:+.3f} GAAP {r['fy27g']:.3f} | FY28 sa {r['fy28_sa']:.3f} "
          f"+ACV {r['fy28']:.3f} | NTM {r['ntm_x']:.3f} Y2 {r['ntm2_x']:.3f} | P/E {r['jpe']:.1f}x | value {r['val']:.2f} | stock {r['ret']:+.1%}")
for lab, key in [("Probability-weighted value ($/sh)", "pw"), ("Expected short return", "esr"), ("Base-case (W2) value ($/sh)", "base"),
                 ("W4 adverse-case short loss (W4 stock return)", "w4loss"), ("Expected short return, EPS-only (no multiple change)", "epsonly")]:
    close(key, g("Valuation", f"B{pick(lab)}"), ref[key])
close("stress", g("Valuation", f"B{pick('Rate-cut stress short loss (stock return)')}"), ref["stress"])
print(f"PW {ref['pw']:.2f} | ESR {ref['esr']:+.1%} | base {ref['base']:.2f} | W4 loss {ref['w4loss']:.1%} | stress {ref['stress']:.1%} "
      f"(value {ref['stress_val']:.2f}) | EPS-only {ref['epsonly']:+.1%}")

# --- sensitivity engine vs Python
T = wb["Sensitivity"]
eng = {T.cell(r, 1).value: r for r in range(1, T.max_row + 1) if isinstance(T.cell(r, 1).value, str) and T.cell(r, 1).value.startswith(("Drag ", "Shares "))}
w2 = deepcopy(mv.WORLDS[1][2])
p0 = mv.BASE
for lab, r in eng.items():
    p = dict(p0)
    w = dict(w2)
    if lab.startswith("Drag"):
        L, F_ = re.findall(r"(\d+)%", lab)
        w["loss"], p["fix"] = -int(L) / 100, int(F_) / 100
    else:
        p["shares"] = float(re.findall(r"(\d+)m", lab)[0])
    o = mv.world(w, p)
    N, t = p["shares"], p["tax"]
    dI = -(p0["shares"] - N) * p["price"] * (p["intq1"] * 4 / (p["cash0"] * 1000)) / 4
    q1 = o["q"]["core"][0] + (p["intq1"] + dI) * (1 - t) / N
    fy = sum(o["q"]["core"]) + 4 * (p["intq1"] + dI) * (1 - t) / N
    val = o["vexc"] + (p["cash_val"] - (p0["shares"] - N) * p["price"] / 1000) * 1000 / N
    close(f"engine {lab} Q1", g("Sensitivity", f"I{r}"), q1)
    close(f"engine {lab} FY27", g("Sensitivity", f"J{r}"), fy)
    close(f"engine {lab} value", g("Sensitivity", f"M{r}"), val)
print("Sensitivity engine cases checked:", len(eng))

# cost of equity shifts / ACV FY28 / close timing
Tr = rows("Sensitivity")
for lab, shf in [("Cost of equity -1.0pt", -0.01), ("Cost of equity -0.5pt", -0.005), ("Cost of equity (base)", 0.0),
                 ("Cost of equity +0.5pt", 0.005), ("Cost of equity +1.0pt", 0.01)]:
    p = dict(p0); p["rf"] += shf
    close(f"CoE {lab}", g("Sensitivity", f"E{Tr[lab]}"), mv.run(p)["esr"])
r_acv = Tr["Expected short return"]
for k, cval in enumerate([-0.04, 0.0, 0.02, 0.04]):
    ws_ = [(n, pr, dict(w, acv28=cval)) for n, pr, w in mv.WORLDS]
    close(f"ACV28 {cval}", g("Sensitivity", f"{'BCDE'[k]}{r_acv}"), mv.run(p0, ws_)["esr"])
    close(f"ACV28 W2 value {cval}", g("Sensitivity", f"{'BCDE'[k]}{Tr['Base-case (W2) value ($/sh)']}"), mv.run(p0, ws_)["base"])
for lab, fr in [("Early Q2 (~Nov 1, 2026)", 1.0), ("Mid Q2 (~Dec 15, 2026; model base)", 0.5), ("Late Q2 / calendar year-end (~Dec 31, 2026)", 1 / 3)]:
    p = dict(p0); p["acv_frac"] = fr
    close(f"close {lab}", g("Sensitivity", f"C{Tr[lab]}"), mv.world(w2, p)["fy27g"])
    print(f"ACV close {lab}: W2 FY27 GAAP {mv.world(w2, p)['fy27g']:.3f}")

# --- checks sheet
X = wb["Checks"]
chk = [(X.cell(r, 1).value, vals.get(("CHECKS", f"B{r}"))) for r in range(11, 20) if X.cell(r, 1).value and isinstance(vals.get(("CHECKS", f"B{r}")), str)]
print("Integrity checks:", chk)
fails += [f"check {a}" for a, b in chk if b != "OK"]
print("Calibration FY26 replay diff: %.4f" % g("Checks", "F7"))

# --- hygiene
fills = [(ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows() for c in row
         if c.fill is not None and c.fill.fill_type not in (None, "none")]
print("Filled cells:", fills[:5] if fills else "NONE", "| hidden sheets:", [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"],
      "| defined names:", [n for n in wb.defined_names if not n.startswith("_xlnm")], "| sheets:", wb.sheetnames)
z = zipfile.ZipFile(path)
print("External links:", [n for n in z.namelist() if "externalLink" in n] or "NONE")
blob = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml") or n.endswith(".vml")).lower()
for term in ["progressive", "90%", "$209", "209m", "jpm", "jpmorgan", "proprietary", "hopkins", "jhu", "university", "arjun", "vivaan", "gupta",
             "singh", "club", "largest online", "post-miss", "what is already priced"]:
    n = blob.count(term)
    if n:
        fails.append(f"banned term '{term}' x{n}")
print("Banned-term scan:", "clean" if not any(f.startswith("banned") for f in fails) else [f for f in fails if f.startswith("banned")])
print("Workbook props:", wb.properties.creator, "|", wb.properties.lastModifiedBy)
print("RESULT:", "ALL MATCH" if not fails else "FAIL", fails[:20])
