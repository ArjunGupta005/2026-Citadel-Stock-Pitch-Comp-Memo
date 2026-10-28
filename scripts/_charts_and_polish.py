# ------------------------------------------------------------------ Summary charts (data linked; placed right of tables)
from math import ceil

from openpyxl.chart.axis import ChartLines
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.chart.title import Title
from openpyxl.chart.text import Text
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun

put(M, "A34", "Chart data: FY27 quarterly EPS ($/sh), linked to Consensus and EPS_Build", font=F_SEC)
for c in range(1, 4):
    M.cell(34, c).border = SEC_LINE
for c, lab in zip("ABC", ["Quarter", "Zacks consensus", "Our ex-ACV estimate"]):
    put(M, f"{c}35", lab, font=F_HDR, border=HDR_LINE, align="left" if c == "A" else "right")
for q in range(4):
    rr = 36 + q
    put(M, f"A{rr}", ["Q1", "Q2", "Q3", "Q4"][q], border=ROW_LINE)
    put(M, f"B{rr}", f"=Consensus!C{5 + q}", USD, F_LINK, align="right", border=ROW_LINE)
    put(M, f"C{rr}", f"=EPS_Build!{get_column_letter(2 + q)}{R2['eps_x']}", USD, F_LINK, align="right", border=ROW_LINE)


def _rich(size=800, bold=False, color="000000"):
    cp = CharacterProperties(latin=DFont(typeface=FN), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _title(text, size=1000):
    cp = CharacterProperties(latin=DFont(typeface=FN), sz=size, b=True, solidFill=NAVY)
    t = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])])))
    t.overlay = False
    return t


def _clean(ch):
    ch.graphical_properties = GraphicalProperties(solidFill="FFFFFF", ln=LineProperties(solidFill="D9D9D9", w=6350))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.txPr = _rich()
    ch.y_axis.txPr = _rich()
    ch.x_axis.majorGridlines = None
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="808080", w=6350))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


# Chart 1: clustered columns, consensus (navy) vs ours (dark red); light horizontal gridlines only
ch1 = BarChart()
ch1.type = "col"
ch1.grouping = "clustered"
ch1.title = _title("FY27 Quarterly EPS: Consensus vs Margin Callers")
ch1.add_data(Reference(M, min_col=2, max_col=3, min_row=35, max_row=39), titles_from_data=True)
ch1.set_categories(Reference(M, min_col=1, min_row=36, max_row=39))
_clean(ch1)
ch1.y_axis.numFmt = "$0.00"
ch1.y_axis.scaling.min = 0
ch1.y_axis.scaling.max = 0.5
ch1.y_axis.majorUnit = 0.1
ch1.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E3E3E3", w=6350)))
ch1.series[0].graphicalProperties = GraphicalProperties(solidFill="17365D", ln=LineProperties(noFill=True))
ch1.series[1].graphicalProperties = GraphicalProperties(solidFill="9B2C2C", ln=LineProperties(noFill=True))
ch1.gapWidth = 80
ch1.legend.position = "b"
ch1.legend.overlay = False
ch1.legend.txPr = _rich()
ch1.height, ch1.width = 7.4, 15.0
M.add_chart(ch1, "J5")

# Chart 2: horizontal bars of scenario value (single series; W1/W2 navy, W3/W4 gray); labels at bar ends
ch2 = BarChart()
ch2.type = "bar"
ch2.grouping = "clustered"
ch2.title = _title("Scenario Value per Share")
ch2.add_data(Reference(M, min_col=7, min_row=20, max_row=23), titles_from_data=False)
ch2.set_categories(Reference(M, min_col=1, min_row=20, max_row=23))
_clean(ch2)
ch2.legend = None
ch2.x_axis.scaling.orientation = "maxMin"          # W1 at top
ch2.y_axis.crosses = "max"                          # value axis at bottom when categories are reversed
ch2.y_axis.scaling.min = 0
ch2.y_axis.scaling.max = 35
ch2.y_axis.majorUnit = 5
ch2.y_axis.numFmt = "$0"
ch2.y_axis.majorGridlines = None
ch2.y_axis.title = _title("Value per share ($); current price $27.29", size=800)
s2 = ch2.series[0]
s2.graphicalProperties = GraphicalProperties(solidFill="17365D", ln=LineProperties(noFill=True))
for idx, col in enumerate(["17365D", "17365D", "8C96A5", "8C96A5"]):
    pt = DataPoint(idx=idx)
    pt.graphicalProperties = GraphicalProperties(solidFill=col, ln=LineProperties(noFill=True))
    s2.dPt.append(pt)
s2.dLbls = DataLabelList()
s2.dLbls.showVal = True
s2.dLbls.showSerName = s2.dLbls.showCatName = s2.dLbls.showLegendKey = s2.dLbls.showPercent = False
s2.dLbls.position = "outEnd"
s2.dLbls.numFmt = "$0.00"
s2.dLbls.txPr = _rich(size=800, bold=True)
ch2.gapWidth = 60
ch2.height, ch2.width = 7.4, 15.0
M.add_chart(ch2, "J21")
M.column_dimensions["I"].width = 3

# ================================================================== global visual pass (styling only)
THIN_ROW = Border(bottom=Side(style="thin", color="D9D9D9"))
FREEZE = {"Summary": None, "Inputs": "A6", "Scenarios": "A6", "EPS_Build": "A4", "Valuation": "A5", "Sensitivity": "A5",
          "Consensus": "A5", "Catalysts": "A5", "Sources": "A4", "Checks": "A4"}
CENTER_COLS = {"Inputs": {"C"}}
RIGHT_TEXT_COLS = {"Consensus": {"B"}, "Sources": {"F"}}


def _width(ws, col_letter):
    return ws.column_dimensions[col_letter].width or 8.43


for ws in wb.worksheets:
    # replace hair rules with thin light-gray rules (borders only; never fills)
    for row in ws.iter_rows():
        for c in row:
            b = c.border
            if b is not None and b.bottom is not None and b.bottom.style == "hair":
                c.border = Border(left=b.left, right=b.right, top=b.top, bottom=Side(style="thin", color="D9D9D9"))
            # alignment: numbers / formulas right, text left (keep explicit centre / wrap)
            if c.value is None:
                continue
            al = c.alignment
            col = c.column_letter
            is_num = isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("="))
            if col in CENTER_COLS.get(ws.title, set()) and c.row > 5:
                horiz = "center"
            elif col in RIGHT_TEXT_COLS.get(ws.title, set()) and c.row > 4:
                horiz = "right"
            elif al.horizontal in ("center", "right") and not is_num:
                horiz = al.horizontal
            else:
                horiz = "right" if is_num else "left"
            c.alignment = Alignment(horizontal=horiz, vertical=al.vertical or ("top" if al.wrap_text else "center"),
                                    wrap_text=al.wrap_text)
    # row heights: fit wrapped text (Arial 9 ≈ 1.15 characters per width unit, 12pt per line)
    merged = {}
    for rng in ws.merged_cells.ranges:
        w = sum(_width(ws, get_column_letter(cc)) for cc in range(rng.min_col, rng.max_col + 1))
        merged[(rng.min_row, rng.min_col)] = (w, rng.max_row - rng.min_row + 1)
    for r in range(1, ws.max_row + 1):
        lines_needed = 1
        for cc in range(1, ws.max_column + 1):
            cell = ws.cell(r, cc)
            if isinstance(cell.value, str) and cell.alignment.wrap_text and not cell.value.startswith("="):
                w, nrows = merged.get((r, cc), (_width(ws, get_column_letter(cc)), 1))
                lines_needed = max(lines_needed, ceil(ceil(len(cell.value) / (w * 1.15)) / nrows))
        if lines_needed > 1:
            ws.row_dimensions[r].height = 12.5 * lines_needed + 3
        elif ws.row_dimensions[r].height and ws.row_dimensions[r].height > 16 and r > 3:
            ws.row_dimensions[r].height = None
    # panes: top-row freezes only; no split panes, no frozen columns
    ws.sheet_view.pane = None
    ws.freeze_panes = FREEZE[ws.title]
    ws.sheet_view.selection[0].activeCell = "A1"
    ws.sheet_view.selection[0].sqref = "A1"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    # print area = used model area only
    last_col = "H" if ws.title == "Summary" else get_column_letter(ws.max_column)
    ws.print_area = f"A1:{last_col}{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

# Summary print area includes the charts (J:R)
M.print_area = f"A1:R{M.max_row}"
wb.active = 0

out = ROOT / "final" / "Margin_Callers.xlsx"
wb.save(out)
print("saved", out)
